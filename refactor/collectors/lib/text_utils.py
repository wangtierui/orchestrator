# -*- coding: utf-8 -*-
"""本地文本/文档工具（移植自 std_lib.scraper_std.crawler_common 第 2~5 节）。

纯标准库 + 惰性第三方导入；零硬依赖。多格式文档抽取（PDF/docx/xlsx/ole2）透明标记
sha256 与抽取状态；归一化/哈希工具；标准化附件记录构造。OCR 调用改走本地 lib.ocr。
"""
import hashlib
import logging
import os
import random
import re
import time
import urllib.error
import urllib.request
import zipfile

LOG = logging.getLogger("refactor.collectors.lib.text_utils")

_ILLEGAL = re.compile(r'[\\/:*?"<>|\r\n\t]+')
_ATTACH_EXT = (".pdf", ".doc", ".docx", ".xls", ".xlsx", ".wps", ".ceb", ".rtf", ".zip", ".rar")


def is_attachment_url(url: str, text: str = "") -> bool:
    """判断链接是否指向文档下载（扩展名命中或链接文案含下载关键词）。"""
    if not url:
        return False
    low = url.split("?")[0].lower()
    if low.endswith(_ATTACH_EXT):
        return True
    kw = ("下载", "download", "word", "pdf", "全文", "附件", "doc", "xlsx", "excel")
    return any(k in (text or "").lower() for k in kw)


def sniff_kind(data: bytes, name: str = "") -> str:
    """依据文件头魔数纠正真实格式。返回 pdf/docx/xlsx/ole2/zip/rar/unknown/empty。"""
    if not data:
        return "empty"
    head = data[:8]
    if head[:4] == b"%PDF":
        return "pdf"
    if head[:4] in (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08"):
        if _is_office_openxml(data):
            if b"xl/" in data[: min(len(data), 2_000_000)] or b"workbook" in data[: min(len(data), 2_000_000)]:
                return "xlsx"
            return "docx"
        ext = os.path.splitext(name)[1].lower()
        if ext == ".docx":
            return "docx"
        if ext == ".xlsx":
            return "xlsx"
        return "zip"
    if head[:8] in (b"Rar!\x1a\x07",) or head[:4] == b"Rar!":
        return "rar"
    if head[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return "ole2"
    ext = os.path.splitext(name)[1].lower()
    if ext in (".pdf",):
        return "pdf"
    if ext in (".docx",):
        return "docx"
    if ext in (".xlsx",):
        return "xlsx"
    if ext in (".xls", ".doc", ".wps", ".ceb", ".rtf"):
        return "ole2"
    return "unknown"


def _is_office_openxml(data: bytes) -> bool:
    try:
        if not data.startswith(b"PK"):
            return False
        with zipfile.ZipFile(__import__("io").BytesIO(data[: min(len(data), 5_000_000)])) as z:
            names = z.namelist()
        return "[Content_Types].xml" in names
    except Exception:  # noqa: BLE001
        return "docProps" in data[: min(len(data), 200_000)].decode("latin-1", "ignore")


def safe_filename(name: str, url: str = "", max_len: int = 120) -> str:
    """生成安全文件名：去除非法字符、限长、保留可读标题 + URL 哈希前缀去重。"""
    raw = (name or "").strip()
    raw = _ILLEGAL.sub("_", raw)
    raw = raw.strip("._ ")
    if not raw:
        raw = "attachment"
    if len(raw) > max_len:
        raw = raw[:max_len].rstrip("._ ")
    if url:
        prefix = hashlib.md5(url.encode("utf-8")).hexdigest()[:8]
        return f"{prefix}_{raw}"
    return raw


def sha256_of(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def garble_ratio(text: str) -> float:
    if not text:
        return 0.0
    return text.count("\ufffd") / max(1, len(text))


def extract_document_text(data: bytes, name: str = "", *, enable_ocr: bool = False, ocr_timeout: int = 60) -> dict:
    """抽取文档内部文本，返回结构化结果。

    {text, kind, extracted, extract_status, sha256, needs_ocr, garble_ratio, size_bytes}
    """
    kind = sniff_kind(data, name)
    sha = sha256_of(data)
    rec = {
        "text": "", "kind": kind, "extracted": False, "extract_status": "unsupported",
        "sha256": sha, "needs_ocr": False, "garble_ratio": 0.0, "size_bytes": len(data),
    }
    if kind == "empty":
        rec["extract_status"] = "empty"
        return rec
    if kind in ("zip", "rar", "unknown"):
        rec["extract_status"] = "unsupported"
        return rec
    try:
        if kind == "pdf":
            return _merge(rec, _extract_pdf(data, enable_ocr, ocr_timeout))
        if kind == "docx":
            return _merge(rec, _extract_docx(data))
        if kind == "xlsx":
            return _merge(rec, _extract_xlsx(data))
        if kind == "ole2":
            return _merge(rec, _extract_ole2(data))
    except Exception as e:  # noqa: BLE001
        rec["extract_status"] = "corrupt"
        rec["text"] = ""
        LOG.warning("抽取 %s(%s) 异常：%s", name, kind, e)
        return rec
    rec["extract_status"] = "unsupported"
    return rec


def _merge(rec: dict, sub: dict) -> dict:
    rec.update(sub)
    rec["garble_ratio"] = garble_ratio(sub.get("text", ""))
    return rec


_PAGENO_RE = re.compile(r"^\s*[—\-–一]\s*\d+\s*[—\-–]?\s*$|^\s*第\s*\d+\s*页(共\d+页)?\s*$|^\s*\d{1,3}\s*$")


def clean_pdf_text(text: str, min_repeat: int = 3) -> str:
    """PDF 文本页眉/页脚清理：按 \f 分页统计行频，跨页重复短行与页码行删除。"""
    if not text:
        return text
    pages = text.split("\f")
    if len(pages) <= 1:
        pages = [text]
    freq: dict[str, int] = {}
    parsed = []
    for pg in pages:
        lines = [ln.strip() for ln in pg.splitlines()]
        body = [ln for ln in lines if ln]
        parsed.append(body)
        for ln in body[:3] + body[-3:]:
            if len(ln) <= 40:
                freq[ln] = freq.get(ln, 0) + 1
    out = []
    for body in parsed:
        keep = []
        for ln in body:
            if _PAGENO_RE.match(ln):
                continue
            if freq.get(ln, 0) >= min_repeat and len(ln) <= 40:
                continue
            keep.append(ln)
        out.append("\n".join(keep))
    return "\f".join(out).strip()


_TEXT_LAYER_MIN_CJK = 30
_EMPTY_QUOTE_RE = re.compile(r'[“"][”"]')


def cjk_count(s: str) -> int:
    return len(re.findall(r"[\u4e00-\u9fa5]", s or ""))


def text_layer_ok(t: str, *, min_cjk: int = _TEXT_LAYER_MIN_CJK) -> bool:
    return cjk_count(t) >= min_cjk and not _EMPTY_QUOTE_RE.search(t)


def has_extraction_gap(t: str) -> bool:
    return bool(_EMPTY_QUOTE_RE.search(t))


def _extract_pdf(data: bytes, enable_ocr: bool, ocr_timeout: int) -> dict:
    cands: list[str] = []
    try:
        import pymupdf

        with pymupdf.open(stream=data, filetype="pdf") as doc:
            t = clean_pdf_text("".join(p.get_text() for p in doc))
        if t.strip():
            if text_layer_ok(t):
                return {"text": t, "extracted": True, "extract_status": "ok", "needs_ocr": False}
            cands.append(t)
    except Exception:  # noqa: BLE001
        pass
    try:
        from pypdf import PdfReader

        reader = PdfReader(__import__("io").BytesIO(data))
        text = ""
        for page in reader.pages:
            try:
                text += (page.extract_text() or "") + "\n"
            except Exception:  # noqa: BLE001
                pass
        t = clean_pdf_text(text)
        if t.strip():
            if text_layer_ok(t):
                return {"text": t, "extracted": True, "extract_status": "ok", "needs_ocr": False}
            cands.append(t)
    except ImportError:
        pass
    try:
        import pdfplumber

        text = ""
        with pdfplumber.open(__import__("io").BytesIO(data)) as pdf:
            for page in pdf.pages:
                try:
                    text += (page.extract_text() or "") + "\n"
                except Exception:  # noqa: BLE001
                    pass
        t = clean_pdf_text(text)
        if t.strip():
            if text_layer_ok(t):
                return {"text": t, "extracted": True, "extract_status": "ok", "needs_ocr": False}
            cands.append(t)
    except ImportError:
        pass
    best = max(cands, key=cjk_count) if cands else ""
    if not _pdf_lib_available():
        return {"text": "", "extracted": False, "extract_status": "library_missing", "needs_ocr": False}
    if enable_ocr and _ocr_available():
        try:
            ocr_text, ocr_engine = _run_ocr_detail(data, ocr_timeout)
            if ocr_text.strip():
                return {"text": ocr_text.strip(), "extracted": True, "extract_status": "ok",
                        "needs_ocr": True, "ocr_engine": ocr_engine}
        except Exception as e:  # noqa: BLE001
            LOG.warning("PDF OCR 失败：%s", e)
    if best.strip():
        return {"text": best, "extracted": True, "extract_status": "ok", "needs_ocr": True}
    return {"text": "", "extracted": False, "extract_status": "needs_ocr", "needs_ocr": True}


def _pdf_lib_available() -> bool:
    try:
        import pymupdf  # noqa: F401
        return True
    except ImportError:
        pass
    try:
        import pypdf  # noqa: F401
        return True
    except ImportError:
        pass
    try:
        import pdfplumber  # noqa: F401
        return True
    except ImportError:
        return False


def _extract_docx(data: bytes) -> dict:
    try:
        import docx

        d = docx.Document(__import__("io").BytesIO(data))
        text = "\n".join(p.text for p in d.paragraphs if p.text)
        if text.strip():
            return {"text": text.strip(), "extracted": True, "extract_status": "ok"}
    except ImportError:
        pass
    try:
        with zipfile.ZipFile(__import__("io").BytesIO(data)) as z:
            xml = z.read("word/document.xml").decode("utf-8", "replace")
        texts = re.findall(r"<w:t[^>]*>(.*?)</w:t>", xml, re.S)
        text = "".join(texts)
        if text.strip():
            return {"text": re.sub(r"\s+", " ", text).strip(), "extracted": True, "extract_status": "ok"}
    except Exception:  # noqa: BLE001
        pass
    if not _docx_lib_available():
        return {"text": "", "extracted": False, "extract_status": "library_missing"}
    return {"text": "", "extracted": False, "extract_status": "empty"}


def _docx_lib_available() -> bool:
    try:
        import docx  # noqa: F401
        return True
    except ImportError:
        return False


def _extract_xlsx(data: bytes) -> dict:
    try:
        import openpyxl

        wb = openpyxl.load_workbook(__import__("io").BytesIO(data), data_only=True, read_only=True)
        rows = []
        for ws in wb.worksheets:
            for r in ws.iter_rows(values_only=True):
                cells = ["" if c is None else str(c) for c in r]
                rows.append("\t".join(cells))
        text = "\n".join(rows)
        if text.strip():
            return {"text": text.strip(), "extracted": True, "extract_status": "ok"}
    except ImportError:
        pass
    try:
        with zipfile.ZipFile(__import__("io").BytesIO(data)) as z:
            names = z.namelist()
            if "xl/sharedStrings.xml" in names:
                xml = z.read("xl/sharedStrings.xml").decode("utf-8", "replace")
                strings = re.findall(r"<t[^>]*>(.*?)</t>", xml, re.S)
                text = "\n".join(strings)
                if text.strip():
                    return {"text": re.sub(r"\s+", " ", text).strip(), "extracted": True, "extract_status": "ok"}
    except Exception:  # noqa: BLE001
        pass
    if not _xlsx_lib_available():
        return {"text": "", "extracted": False, "extract_status": "library_missing"}
    return {"text": "", "extracted": False, "extract_status": "empty"}


def _xlsx_lib_available() -> bool:
    try:
        import openpyxl  # noqa: F401
        return True
    except ImportError:
        return False


_WPS_AVAILABLE: bool | None = None


def wps_available() -> bool:
    global _WPS_AVAILABLE
    if _WPS_AVAILABLE is not None:
        return _WPS_AVAILABLE
    _WPS_AVAILABLE = False
    try:
        import winreg

        for prog in ("KWPS.Application", "WPS.Application", "ET.Application"):
            try:
                k = winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, prog + r"\CLSID")
                clsid, _ = winreg.QueryValueEx(k, "")
                winreg.CloseKey(k)
            except OSError:
                continue
            exe = ""
            for view in ("LocalServer32", "InprocServer32"):
                try:
                    k2 = winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, rf"CLSID\{clsid}\{view}")
                    raw, _ = winreg.QueryValueEx(k2, "")
                    winreg.CloseKey(k2)
                    exe = str(raw).strip().strip('"')
                    break
                except OSError:
                    continue
            if exe and os.path.exists(exe):
                _WPS_AVAILABLE = True
                break
    except Exception:  # noqa: BLE001
        _WPS_AVAILABLE = False
    return _WPS_AVAILABLE


def _extract_doc_via_wps(data: bytes, timeout: float = 45.0) -> str | None:
    if not wps_available():
        return None
    import subprocess
    import tempfile
    import threading

    def _wps_pids() -> set:
        try:
            out = subprocess.run(["tasklist", "/FI", "IMAGENAME eq wps.exe", "/FO", "CSV", "/NH"],
                                 capture_output=True, timeout=10)
            raw = out.stdout
            try:
                text = raw.decode("gbk", errors="replace")
            except Exception:  # noqa: BLE001
                text = raw.decode("utf-8", errors="replace")
            pids = set()
            for line in text.strip().splitlines():
                if "wps.exe" not in line:
                    continue
                parts = line.split('","')
                if len(parts) > 1:
                    pids.add(parts[1].strip('"'))
            return pids
        except Exception:  # noqa: BLE001
            return set()

    tmp = ""
    before = _wps_pids()
    result: dict = {"text": None, "done": False}
    try:
        fd, tmp = tempfile.mkstemp(suffix=".doc", prefix="wps_tmp_")
        with os.fdopen(fd, "wb") as f:
            f.write(data)

        def _watchdog() -> None:
            deadline = time.monotonic() + timeout
            while time.monotonic() < deadline and not result["done"]:
                time.sleep(0.3)
            if not result["done"]:
                after = _wps_pids()
                for pid in after - before:
                    try:
                        subprocess.run(["taskkill", "/F", "/PID", pid], capture_output=True, timeout=10)
                    except Exception:  # noqa: BLE001
                        pass

        wt = threading.Thread(target=_watchdog, daemon=True)
        wt.start()
        import win32com.client

        app = win32com.client.Dispatch("KWPS.Application")
        try:
            app.Visible = False
        except Exception:  # noqa: BLE001
            pass
        try:
            app.DisplayAlerts = 0
        except Exception:  # noqa: BLE001
            pass
        try:
            doc = app.Documents.Open(tmp, ReadOnly=True, AddToRecentFiles=False)
            try:
                text = doc.Content.Text or ""
            finally:
                try:
                    doc.Close(False)
                except Exception:  # noqa: BLE001
                    pass
        finally:
            try:
                app.Quit()
            except Exception:  # noqa: BLE001
                pass
        text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
        text = re.sub(r"\r\n?", "\n", text)
        text = re.sub(r"[ \t]{2,}", " ", text)
        text = text.strip()
        result["text"] = text or None
        return result["text"]
    except Exception:  # noqa: BLE001
        return None
    finally:
        result["done"] = True
        if tmp and os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass


def _extract_ole2(data: bytes) -> dict:
    if (b"Workbook" in data[:200_000] or b"\x00B\x00o\x00o\x00k" in data[:4096] or b"Book" in data[:512]):
        try:
            import xlrd

            bk = xlrd.open_workbook(__import__("io").BytesIO(data))
            rows = []
            for sh in bk.sheets():
                for r in range(sh.nrows):
                    rows.append("\t".join(str(c.value) for c in sh.row(r)))
            text = "\n".join(rows)
            if text.strip():
                return {"text": text.strip(), "extracted": True, "extract_status": "ok"}
        except ImportError:
            return {"text": "", "extracted": False, "extract_status": "library_missing"}
        except Exception:  # noqa: BLE001
            pass
    wps_text = _extract_doc_via_wps(data)
    if wps_text:
        return {"text": wps_text, "extracted": True, "extract_status": "ok_wps_com"}
    try:
        import olefile

        if not olefile.isOleFile(__import__("io").BytesIO(data)):
            return {"text": "", "extracted": False, "extract_status": "unsupported"}
        ol = olefile.OleFileIO(__import__("io").BytesIO(data))
        text_parts = []
        for stream in ol.listdir():
            try:
                raw = ol.openstream(stream).read()
                text_parts.append(raw.decode("utf-16-le", "ignore"))
            except Exception:  # noqa: BLE001
                pass
        text = "\n".join(text_parts)
        text = re.sub(r"[\x00-\x1f\x7f-\x9f]", " ", text)
        text = re.sub(r"\s+", " ", text).strip()
        if text:
            return {"text": text, "extracted": True, "extract_status": "ok"}
        return {"text": "", "extracted": False, "extract_status": "empty"}
    except ImportError:
        return {"text": "", "extracted": False, "extract_status": "library_missing"}
    except Exception:  # noqa: BLE001
        return {"text": "", "extracted": False, "extract_status": "corrupt"}


# ---- OCR 桥接（本地 lib.ocr，去掉 config.loader / ocr_correction 外部依赖） -------
def _ocr_available() -> bool:
    try:
        from .ocr import get_ocr
    except Exception as e:  # noqa: BLE001
        LOG.warning("本地 OCR 模块不可用：%s", e)
        return False
    try:
        return any(get_ocr().health().values())
    except Exception as e:  # noqa: BLE001
        LOG.warning("本地 OCR 不可用：%s", e)
        return False


def _run_ocr(data: bytes, timeout: int) -> str:
    from .ocr import get_ocr
    import os
    import tempfile

    fd, tmp = tempfile.mkstemp(suffix=".pdf", prefix="ocr_run_")
    try:
        with os.fdopen(fd, "wb") as f:
            f.write(data)
        res = get_ocr().extract_pdf(tmp, force_ocr=True)
        return res.text or ""
    except Exception as e:
        raise RuntimeError(f"OCR 执行失败: {e}") from e
    finally:
        try:
            os.remove(tmp)
        except OSError:
            pass


def _run_ocr_detail(data: bytes, timeout: int) -> tuple[str, str]:
    from .ocr import get_ocr
    import os
    import tempfile

    fd, tmp = tempfile.mkstemp(suffix=".pdf", prefix="ocr_run_")
    try:
        with os.fdopen(fd, "wb") as f:
            f.write(data)
        res = get_ocr().extract_pdf(tmp, force_ocr=True)
        return (res.text or ""), (getattr(res, "engine", "") or "")
    except Exception as e:
        raise RuntimeError(f"OCR 执行失败: {e}") from e
    finally:
        try:
            os.remove(tmp)
        except OSError:
            pass


# --------------------------------------------------------------------------- #
# 4. 归一化 / 哈希工具
# --------------------------------------------------------------------------- #
_DATE_RE1 = re.compile(r"(\d{4})[-年](\d{1,2})[-月](\d{1,2})日?")
_DATE_RE2 = re.compile(r"(\d{4})[-/](\d{1,2})[-/](\d{1,2})")


def normalize_date(text: str) -> str:
    if not isinstance(text, str) or not text:
        return ""
    m = _DATE_RE1.search(text) or _DATE_RE2.search(text)
    if not m:
        return ""
    y, mo, d = m.groups()
    try:
        return "%04d-%02d-%02d" % (int(y), int(mo), int(d))
    except ValueError:
        return ""


_FULLWIDTH = str.maketrans("０１２３４５６７８９", "0123456789")


def normalize_digits(s: str) -> str:
    if not s:
        return ""
    return str(s).translate(_FULLWIDTH)


def content_sha256(text: str) -> str:
    return sha256_of((text or "").encode("utf-8"))


def build_attachment_record(
    file_name: str, file_url: str, local_path: str, data: bytes, *,
    entry_id: str = "", text: str = "", extracted: bool = False,
    extract_status: str = "unsupported", needs_ocr: bool = False,
) -> dict:
    """构造标准化附件记录（与 nfra 质量基线一致）。"""
    kind = sniff_kind(data, file_name)
    return {
        "entry_id": entry_id,
        "file_name": file_name,
        "file_url": file_url,
        "local_path": local_path,
        "size_bytes": len(data),
        "extension": os.path.splitext(file_name)[1].lower(),
        "kind": kind,
        "sha256": sha256_of(data),
        "extracted": extracted,
        "extract_status": extract_status,
        "needs_ocr": needs_ocr,
        "garble_ratio": garble_ratio(text),
        "text": text,
        "fetch_status": "ok",
    }
