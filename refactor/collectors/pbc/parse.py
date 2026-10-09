# -*- coding: utf-8 -*-
"""pbc 解析/抽取工具（原 pbc_parse.py 迁出，逻辑逐函数等价）。

含：文本清洗、文件名/URL 安全化、附件格式探测与多格式正文抽取，
以及详情页结构化字段解析（parse_detail / extract_doc_number / extract_issuing / extract_effective）。
"""
import html
import os
import re
import sys
import urllib.parse
import urllib.request

# ---- 仓库引导：使 std_lib 可导入（orchestrator 根）----
_ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", ".."))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)
del _ROOT

from std_lib.scraper_std.doc_convert import find_libreoffice

from .net import (
    ATTACH_EXT,
    ISSUING_AUTHORITIES,
    RE_ARTICLE_TITLE_META,
    RE_DOC_NUMBER,
    RE_EFFECTIVE,
    RE_P,
    RE_SHIJIAN,
    RE_TAG,
    RE_TITLE_H2,
    RE_WS,
    RE_ZOOM,
)

LO_PATH = find_libreoffice()
LO_AVAILABLE = LO_PATH is not None


def safe_filename(title, ext, url):
    """生成本地安全文件名：hash 前缀防重名 + 可读标题 + 小写扩展名。

    核心逻辑（去非法字符 / 限长 / URL 哈希前缀）委托 crawler_common.safe_filename，
    本函数仅补扩展名；保留 pbc 原有"限长 60"，签名与返回结构不变。
    """
    repo_root = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", ".."))
    if repo_root not in sys.path:
        sys.path.insert(0, repo_root)
    from std_lib.scraper_std.crawler_common import safe_filename as _cc_safe_filename
    base = _cc_safe_filename(title, url, 60)
    ext = (ext or "bin").lower()
    return f"{base}.{ext}"


def detect_magic(bdata):
    """按文件头 magic bytes 判断实际格式，处理扩展名与内容不符的'格式不匹配'异常。"""
    if not bdata:
        return None
    if bdata[:4] == b"%PDF":
        return "pdf"
    if bdata[:4] in (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08"):
        return "zip"          # docx / xlsx 均为 zip 容器
    if bdata[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return "ole"          # .doc / 旧 .xls 为 OLE 复合文档
    return None


def detect_libreoffice():
    """兼容别名：委托共享 find_libreoffice（五源统一探测）。"""
    return find_libreoffice()


def convert_with_libreoffice(doc_path):
    """兼容别名：委托共享 doc_to_docx（.doc/.wps/.rtf/.ceb → .docx，headless）。不可用返回 None。"""
    from std_lib.scraper_std.doc_convert import doc_to_docx
    return doc_to_docx(doc_path)


def extract_pdf_text(data):
    """用统一 OCR 模块抽取 PDF 文本（文本层优先，扫描件走 OCR 引擎链）。"""
    import tempfile
    repo_root = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", ".."))
    if repo_root not in sys.path:
        sys.path.insert(0, repo_root)
    from std_lib.scraper_std.ocr_engine import get_ocr
    fd, tmp = tempfile.mkstemp(suffix=".pdf", prefix="pbc_main_ocr_")
    try:
        with os.fdopen(fd, "wb") as f:
            f.write(data)
        res = get_ocr().extract_pdf(tmp, force_ocr=False)
    finally:
        try:
            os.remove(tmp)
        except OSError:
            pass
    if res.success and res.text.strip():
        return res.text
    return None


def extract_xls_text(data):
    try:
        import io

        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        out = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    out.append(" | ".join(cells))
        return "\n".join(out).strip()
    except Exception:  # noqa: BLE001
        return None


def strip_tags(s):
    if not s:
        return ""
    s = RE_TAG.sub("", s)
    s = html.unescape(s)
    return RE_WS.sub(" ", s).strip()


def clean_text(s):
    return RE_WS.sub(" ", html.unescape(s or "")).strip()


def normalize_digits(s):
    """将全角数字 ０-９ 归一为半角 0-9，统一日期口径。"""
    if not s:
        return s
    trans = str.maketrans("０１２３４５６７８９", "0123456789")
    return s.translate(trans)


def is_attachment(url):
    low = url.lower().split("?")[0]
    return any(low.endswith(ext) for ext in ATTACH_EXT)


def safe_url(url):
    """
    对非 ASCII 字符（如中文文件名）的 URL 路径/查询做百分号编码，
    避免 urllib 抛出 'unknown url type' / 'ascii codec' 错误。
    仅编码 path 与 query，保留 scheme/netloc 及 '/'、'=&' 等分隔符。
    """
    parts = urllib.parse.urlsplit(url)
    path = urllib.parse.quote(parts.path, safe="/")
    query = urllib.parse.quote(parts.query, safe="=&")
    return urllib.parse.urlunsplit((parts.scheme, parts.netloc, path, query, parts.fragment))


def extract_docx_text(data):
    """尝试用 python-docx 提取 .docx 文本；不可用则返回 None。"""
    try:
        from docx import Document
    except Exception:  # noqa: BLE001
        return None
    import io
    try:
        doc = Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs if p.text).strip()
    except Exception:  # noqa: BLE001
        return None


# ----------------------------------------------------------------------------------
# 详情页结构化字段解析
# ----------------------------------------------------------------------------------
def parse_detail(html_text):
    """从 HTML 详情页提取结构化字段。"""
    # 标题
    title = ""
    mt = RE_ARTICLE_TITLE_META.search(html_text)
    if mt:
        title = clean_text(mt.group(1))
    if not title:
        mh = RE_TITLE_H2.search(html_text)
        if mh:
            title = clean_text(strip_tags(mh.group(1)))

    # 发布日期
    publish_date = None
    ms = RE_SHIJIAN.search(html_text)
    if ms:
        raw = ms.group(1).strip()
        dm = re.search(r"\d{4}[-/年]\d{1,2}[-/月]\d{1,2}", raw)
        if dm:
            publish_date = normalize_digits(dm.group(0)).replace("年", "-").replace("月", "-").replace("/", "-")

    # 正文：优先 #zoom，回退到 <td class="content">
    body_text = ""
    zm = RE_ZOOM.search(html_text)
    if zm:
        paras = [strip_tags(p) for p in RE_P.findall(zm.group(1))]
        body_text = "\n".join(p for p in paras if p).strip()
    if not body_text:
        # 兜底：尝试常见 gov CMS 正文容器（仅当上述选择器均落空时启用，
        # 取含多个 <p> 段落的最长候选，避免抓取导航/页眉等短文本）。
        for token in ("TRS_Editor", "content", "pages_content", "article",
                      "fontarea", "zoom", "newscontent"):
            m = re.search(
                r'<(div|font|td|article|span)[^>]*\bclass="[^"]*%s[^"]*"[^>]*>(.*?)</\1>'
                % re.escape(token), html_text, re.IGNORECASE | re.DOTALL)
            if not m:
                m = re.search(
                    r'<(div|font|td|article|span)[^>]*\bid="%s"[^>]*>(.*?)</\1>'
                    % re.escape(token), html_text, re.IGNORECASE | re.DOTALL)
            if m:
                paras = [strip_tags(p) for p in RE_P.findall(m.group(2))]
                cand = "\n".join(p for p in paras if p).strip()
                if len(cand) > len(body_text):
                    body_text = cand

    # 启发式提取：发文字号 / 发文机关 / 生效日期
    # 注意：仅基于正文 body_text 抽取，避免页面页眉/导航中的"中国人民银行"等干扰。
    doc_number = extract_doc_number(body_text)
    issuing = extract_issuing(body_text)
    effective = extract_effective(body_text)

    return {
        "title": title,
        "publish_date": publish_date,
        "document_number": doc_number,
        "issuing_authority": issuing,
        "effective_date": effective,
        "content": body_text,
    }


def extract_doc_number(*texts):
    """
    启发式提取发文字号。仅匹配高置信模式（含年份六角括号 〔YYYY〕 或机关关键词+号），
    刻意排除纯数字串（如页面追踪 ID '05073439号'），避免产生虚假发文字号。
    注意：国家法律类等由人大通过的文本通常无"发文字号"，此时返回 None 属正常。
    """
    blob = "\n".join(t for t in texts if t)
    # 委托五源统一模块 doc_number（15+ 优先级正则 + 规范化）；
    # 模块不可用/未命中时回退原 RE_DOC_NUMBER，并同样做规范化输出。
    try:
        try:
            from std_lib.scraper_std.doc_number import extract_doc_number as _u
            from std_lib.scraper_std.doc_number import normalize_doc_number as _n
        except ImportError:
            from std_lib.scraper_std.doc_number import extract_doc_number as _u
            from std_lib.scraper_std.doc_number import normalize_doc_number as _n
    except Exception:  # pragma: no cover  # noqa: BLE001
        _u = _n = None  # type: ignore[assignment]
    if _u is not None:
        dn = _u(blob)
        if dn:
            return dn
    m = RE_DOC_NUMBER.search(blob)
    if m:
        return _n(m.group(0)) if _n is not None else clean_text(m.group(0))   # N-67：显式非 None
    return None


def extract_issuing(*texts):
    """
    启发式识别发文机关。优先在前言/标题区（前 350 字）匹配，
    并优先识别"X令"形式（部门规章/规范性文件典型）；法律类前言含
    "全国人民代表大会常务委员会"，不会被正文中出现的"中国人民银行"误判。
    """
    blob = "\n".join(t for t in texts if t)
    head = blob[:350]
    m = re.search(
        r'(中国人民银行|国务院|国家外汇管理局|财政部|国家金融监督管理总局|'
        r'中国银行保险监督管理委员会|中国保险监督管理委员会|'
        r'中国证券监督管理委员会|全国人民代表大会常务委员会|全国人民代表大会)\s*令', head)
    if m:
        return m.group(1)
    for auth in ISSUING_AUTHORITIES:
        if auth in head:
            return auth
    for auth in ISSUING_AUTHORITIES:
        if auth in blob:
            return auth
    return None


def extract_effective(body_text):
    m = RE_EFFECTIVE.search(body_text)
    if m:
        return normalize_digits(clean_text(m.group(1))).replace(" ", "")
    return None


__all__ = [
    "clean_text", "convert_with_libreoffice", "detect_libreoffice", "detect_magic",
    "extract_docx_text", "extract_pdf_text", "extract_xls_text", "is_attachment",
    "normalize_digits", "safe_filename", "safe_url", "strip_tags",
    "parse_detail", "extract_doc_number", "extract_issuing", "extract_effective",
]
