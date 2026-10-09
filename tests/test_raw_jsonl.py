# -*- coding: utf-8 -*-
"""tests.test_raw_jsonl — 批 47（S-A/S-B/N-206）守卫：raw 主库 **JSONL 化** 与**流式**读写。

覆盖：
  ① `pipeline.load_raw_records`：`.jsonl`（首行 `_meta` 跳过）与 `.json`（数组/封装对象）
     **两形态等价**（同一记录集），且 `.json` 走**流式**（不调用 `json.load` 整表）；
  ② `gov_parse.write_outputs`：写 JSONL（首行 `_meta`、一行一记录）+ 与 `MasterView` 读回一致；
  ③ `gov_parse.merge_and_write`：**流式合并**（旧库逐条 → 同键新覆盖 → 新发现追加）语义与
     `merge_with_master` **逐键等价**（同一输入的合并结果集合相同）；
  ④ 迁移兼容：仅有旧 `.json` 时读取侧仍可工作（`_find_master` 回退）。
"""
from __future__ import annotations

import importlib.util
import json
import os
import sys

import pytest

import paths

ROOT = paths.ROOT


def _load(rel_path: str, name: str):
    spec = importlib.util.spec_from_file_location(name, os.path.join(ROOT, rel_path))
    assert spec and spec.loader
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


@pytest.fixture(scope="module")
def gp():
    return _load(os.path.join("modules", "regulatory_scrapers", "collectors", "gov_parse.py"), "_gp_jsonl")


@pytest.fixture(scope="module")
def pipe():
    """按**包路径**导入：`pipeline.py` 含相对导入（`from .cleaner import …`）
    ⇒ 不能用 `spec_from_file_location` 独立加载（实测 ImportError: attempted relative import）。"""
    from std_lib.scraper_std import pipeline

    return pipeline


RECS = [
    {"detail_url": "u1", "title": "甲", "full_text": "正文一", "source": "xzfgk"},
    {"detail_url": "u2", "title": "乙", "full_text": "", "source": "zhengceku"},
    {"title": "丙（无 URL）", "full_text": "正文三"},
]


def _write_json(tmp_path, recs, wrap=True):
    p = tmp_path / "raw.json"
    payload = {"source": "gov", "category": "行政法规+部门文件", "count": len(recs), "records": recs} \
        if wrap else recs
    p.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return str(p)


def _write_jsonl(tmp_path, recs):
    p = tmp_path / "raw.jsonl"
    with open(p, "w", encoding="utf-8") as f:
        f.write(json.dumps({"_meta": {"source": "gov", "format": "jsonl/1"}}, ensure_ascii=False) + "\n")
        for r in recs:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")
    return str(p)


# --------------------------------------------------------------------------- #
# ① load_raw_records 双格式等价 + .json 流式
# --------------------------------------------------------------------------- #
def test_load_json_and_jsonl_equivalent(pipe, tmp_path) -> None:
    """`.json`（封装对象）与 `.jsonl`（含首行 _meta）必须读到**同一记录集**。"""
    a = pipe.load_raw_records(_write_json(tmp_path, RECS))
    b = pipe.load_raw_records(_write_jsonl(tmp_path, RECS))
    assert a == b == RECS, "双格式记录集必须一致，且 _meta 行不得混入"


def test_load_bare_json_array(pipe, tmp_path) -> None:
    """裸数组形态（旧实现支持）仍须可用。"""
    assert pipe.load_raw_records(_write_json(tmp_path, RECS, wrap=False)) == RECS


def test_load_json_does_not_full_load(pipe, tmp_path, monkeypatch) -> None:
    """`.json` 必须走**流式**（判据：不调用 `json.load` 解析整表）。"""
    p = _write_json(tmp_path, RECS)
    calls = {"n": 0}
    real = json.load

    def _spy(fp, *a, **k):
        calls["n"] += 1
        return real(fp, *a, **k)

    monkeypatch.setattr(pipe.json, "load", _spy)
    assert pipe.load_raw_records(p) == RECS
    assert calls["n"] == 0, "不得整表 json.load"


def test_load_meta_only_jsonl(pipe, tmp_path) -> None:
    """只有 `_meta` 行的 JSONL ⇒ 零记录（不得把信封当记录）。"""
    p = tmp_path / "empty.jsonl"
    p.write_text(json.dumps({"_meta": {"source": "gov"}}, ensure_ascii=False) + "\n", encoding="utf-8")
    assert pipe.load_raw_records(str(p)) == []


# --------------------------------------------------------------------------- #
# ② write_outputs → JSONL 往返
# --------------------------------------------------------------------------- #
def test_write_outputs_writes_jsonl_roundtrip(gp, tmp_path) -> None:
    """`write_outputs` 落 JSONL（首行 _meta），`MasterView` 读回一致；索引随写建立。"""
    cfg = gp.ScrapeConfig(source="gov", base_url="", category="行政法规+部门文件",
                          out_dir=str(tmp_path), resume=True)
    paths_ = gp.write_outputs(RECS, cfg)
    assert paths_["json"].endswith("gov_laws.jsonl")
    head = open(paths_["json"], encoding="utf-8").readline()
    assert json.loads(head)["_meta"]["format"] == "jsonl/1"
    mv = gp.MasterView(str(tmp_path))
    assert mv.records() == RECS
    assert os.path.exists(gp.master_index_path(str(tmp_path)))
    seen, detailed = mv.keys()
    assert ("甲", "u1") in seen and detailed == {"u1"}


# --------------------------------------------------------------------------- #
# ③ merge_and_write 与 merge_with_master 语义等价
# --------------------------------------------------------------------------- #
def test_merge_and_write_matches_merge_with_master(gp, tmp_path) -> None:
    """**同键新覆盖 / 新发现追加 / 旧记录保留**：流式合并结果与整表合并**逐键等价**。"""
    cfg = gp.ScrapeConfig(source="gov", base_url="", category="C", out_dir=str(tmp_path), resume=True)
    gp.write_outputs(RECS, cfg)
    new = [{"detail_url": "u1", "title": "甲（改）", "full_text": "新正文"},
           {"detail_url": "u9", "title": "丁", "full_text": "新增"}]
    # 参照实现（整表合并）
    expect = {str(r.get("detail_url") or r.get("title") or ""): r
              for r in gp.merge_with_master(new, str(tmp_path))}
    st = gp.merge_and_write(new, cfg)
    got = {str(r.get("detail_url") or r.get("title") or ""): r for r in gp.MasterView(str(tmp_path)).records()}
    assert set(got) == set(expect), "键集必须一致（旧记录保留 + 新发现追加）"
    assert got["u1"]["title"] == "甲（改）", "同键必须**新优先**覆盖"
    assert got["u2"]["title"] == "乙" and "u9" in got
    assert st["merged"] == 4 and st["replaced"] == 1


def test_master_view_accepts_legacy_json(gp, tmp_path) -> None:
    """迁移兼容：仅有旧 `gov_laws.json` 时，`MasterView` 走**旧格式流式**读取，功能不受阻。"""
    legacy = tmp_path / "gov_laws.json"
    legacy.write_text(json.dumps({"source": "gov", "records": RECS}, ensure_ascii=False), encoding="utf-8")
    mv = gp.MasterView(str(tmp_path))
    assert mv.records() == RECS
    seen, detailed = mv.keys()
    assert ("乙", "u2") in seen and detailed == {"u1"}
    assert gp._find_master(str(tmp_path)).endswith("gov_laws.json")


# --------------------------------------------------------------------------- #
# ⑤ 批 49（T1）：四源 raw 统一 JSONL 的**读写助手**
# --------------------------------------------------------------------------- #
def test_resolve_raw_path_dual(tmp_path) -> None:
    """历史 `.json` 名 → 实际存在的 `.jsonl`；反向亦成立；两者都不存在则原样返回。"""
    from std_lib.scraper_std import pipeline as pipe

    jl = tmp_path / "mof_laws.jsonl"
    jl.write_text('{"_meta": {"source": "mof"}}\n', encoding="utf-8")
    assert pipe.resolve_raw_path(str(tmp_path / "mof_laws.json")).endswith("mof_laws.jsonl")
    assert pipe.resolve_raw_path(str(jl)).endswith(".jsonl")
    js = tmp_path / "only.json"
    js.write_text("[]", encoding="utf-8")
    assert pipe.resolve_raw_path(str(js)).endswith("only.json"), "旧 .json 仍在时应原样命中"
    assert pipe.resolve_raw_path(str(tmp_path / "none.json")).endswith("none.json")


def test_write_and_read_raw_jsonl_roundtrip(tmp_path) -> None:
    """`write_raw_jsonl`：首行 `_meta` 信封 + 一行一记录 + 原子（无 .tmp 残留）；读回等价。"""
    from std_lib.scraper_std import pipeline as pipe

    recs = [{"a": 1, "title": "甲"}, {"b": [1, 2], "title": "乙"}]
    p = pipe.write_raw_jsonl(str(tmp_path / "mof_laws.json"), recs, source="mof",
                             meta={"count": 2, "errors": []})
    assert p.endswith("mof_laws.jsonl") and not os.path.exists(p + ".tmp")
    head = json.loads(open(p, encoding="utf-8").readline())
    assert head["_meta"]["source"] == "mof" and head["_meta"]["count"] == 2 \
        and head["_meta"]["errors"] == [] and head["_meta"]["format"] == "jsonl/1"
    assert pipe.read_raw_records(str(tmp_path / "mof_laws.json")) == recs, \
        "按历史 .json 名也应读到（助手内部解析）"


def test_read_raw_records_supports_bare_array_and_items(tmp_path) -> None:
    """兼容历史形态：裸数组（pbc/supp 旧态）与 `items` 封装（mof 旧态）。"""
    from std_lib.scraper_std import pipeline as pipe

    bare = tmp_path / "pbc_laws.jsonl"
    bare.write_text('{"_meta": {}}\n{"t": 1}\n{"t": 2}\n', encoding="utf-8")
    assert [r["t"] for r in pipe.read_raw_records(str(bare))] == [1, 2]
    wrapped = tmp_path / "old.json"
    wrapped.write_text(json.dumps({"count": 2, "items": [{"t": 3}, {"t": 4}]}), encoding="utf-8")
    assert [r["t"] for r in pipe.read_raw_records(str(wrapped))] == [3, 4]


# --------------------------------------------------------------------------- #
# ⑥ 批 49（N-208）：XLSX 抽取**按内容裁剪 + 上限**（防"声明维度空网格"爆量）
# --------------------------------------------------------------------------- #
def _mk_xlsx(tmp_path, *, far_cell: bool, rows: int = 3):
    """造一个 xlsx；`far_cell=True` 时在远端列放一个单元格以**抬高声明维度**（复现真实缺陷）。"""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "序号"
    ws["B1"] = "名称"
    for r in range(2, rows + 1):
        ws.cell(row=r, column=1, value=r - 1)
        ws.cell(row=r, column=2, value="药品%s" % (r - 1))
    if far_cell:
        ws.cell(row=rows + 1, column=600, value="")   # 空值但存在 ⇒ 声明维度被抬到 600 列
    p = tmp_path / ("far.xlsx" if far_cell else "plain.xlsx")
    wb.save(p)
    return str(p)


def test_xlsx_extract_trims_declared_dimension(tmp_path) -> None:
    """**回归守卫**：声明维度被抬高（600 列）时，抽取文本仍只含**有值内容**（不得按声明展开）。"""
    from std_lib.scraper_std.crawler_common import extract_document_text

    plain = extract_document_text(open(_mk_xlsx(tmp_path, far_cell=False), "rb").read(), "a.xlsx")
    far = extract_document_text(open(_mk_xlsx(tmp_path, far_cell=True), "rb").read(), "b.xlsx")
    t_far = far.get("text") or ""
    assert far.get("extracted") and "药品1" in t_far
    assert len(t_far) < 200, "远端空单元格不得把文本撑大（实测真实文件曾抽出 71,769,545 字符）"
    assert abs(len(t_far) - len(plain.get("text") or "")) < 40, "与无远端单元格时基本等长"


def test_xlsx_extract_has_char_cap(tmp_path, monkeypatch) -> None:
    """上限生效且**显式标注**（不静默丢）。"""
    from std_lib.scraper_std import crawler_common as cc

    monkeypatch.setattr(cc, "XLSX_MAX_CHARS", 80)
    p = _mk_xlsx(tmp_path, far_cell=False, rows=40)
    ex = cc.extract_document_text(open(p, "rb").read(), "c.xlsx")
    t = ex.get("text") or ""
    assert "表格超限" in t and len(t) < 400, "应截断并带标注"


# --------------------------------------------------------------------------- #
# ④ 契约：RAW_JSON 指向 JSONL
# --------------------------------------------------------------------------- #
def test_raw_json_points_to_jsonl() -> None:
    """链侧事实源口径必须指向 JSONL（否则 `_raw_size`/clean 超时按旧文件取大小）。

    批 49（T1）：**五源全部**统一为 JSONL（gov 批 47 已迁；mof/nfra/pbc/supp 本批迁移）。
    """
    from config.enums import SOURCE_ORDER  # 受控枚举：不得在测试里写源集合字面量（门禁 v3）

    rpr = _load(os.path.join("tools", "run_production_refresh.py"), "_rpr_jsonl")
    for src in SOURCE_ORDER:
        assert rpr.RAW_JSON[src].endswith(".jsonl"), src
    clean = _load(os.path.join("modules", "regulatory_scrapers", "clean",
                               "run_clean_pipeline.py"), "_clean_jsonl")
    for src in SOURCE_ORDER:
        assert clean.RAW_MASTER_NAMES[src].endswith(".jsonl"), src
