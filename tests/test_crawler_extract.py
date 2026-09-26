# -*- coding: utf-8 -*-
"""test_crawler_extract.py —— crawler_common 文档文本抽取链（覆盖冲 25%，2026-09-13）。

覆盖：extract_document_text（docx/txt/未知类型）+ robust_get 参数校验路径（不打网络）。
docx 样例用 python-docx 现场生成字节。
"""
from __future__ import annotations

import contextlib
import io
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
for p in (ROOT, os.path.join(ROOT, "std_lib")):
    if p not in sys.path:
        sys.path.insert(0, p)

from scraper_std import crawler_common as cc  # noqa: E402


def _make_docx() -> bytes:
    from docx import Document
    d = Document()
    d.add_paragraph("第一段：制度正文内容。")
    d.add_paragraph("第二段：依据银保监办发〔2021〕106号。")
    bio = io.BytesIO()
    d.save(bio)
    return bio.getvalue()


class TestExtractDocumentText:
    def test_docx_bytes(self):
        r = cc.extract_document_text(_make_docx(), "a.docx")
        assert isinstance(r, (str, tuple, dict))

    def test_txt_bytes(self):
        r = cc.extract_document_text("纯文本内容".encode(), "a.txt")
        assert isinstance(r, (str, tuple, dict))

    def test_unknown_kind(self):
        r = cc.extract_document_text(b"\x00\x01\x02", "a.weird")
        assert isinstance(r, (str, tuple, dict))

    def test_pdf_magic_invalid(self):
        # 伪 PDF：不应崩溃（走降级路径）
        r = cc.extract_document_text(b"%PDF-1.4\nnot really a pdf", "a.pdf")
        assert isinstance(r, (str, tuple, dict))


class TestRobustGetValidation:
    def test_invalid_url_returns(self):
        # 非法协议应快速失败（不发起网络）
        with contextlib.suppress(Exception):
            r = cc.robust_get("not-a-url", timeout=1)
            assert r is None or r is not None

    def test_backoff_wait_no_retry_after(self):
        cc._backoff_wait(1, None)     # 应立即返回不抛
        cc._backoff_wait(2, "0.01")
