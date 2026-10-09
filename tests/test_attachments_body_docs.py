# -*- coding: utf-8 -*-
"""批 50：写入侧「正文载体 vs 真实附件」分流 + XLSX 抽取共享原语的回归守卫。

背景（用户 2026-10-09 命题）：五源写入 raw 时 `attachments` 混入大量**并非公告附件**的条目
—— 它们是"正文以 doc/pdf/word 形式发布"的**正文载体**（gov「下载Word/PDF」、mof fileType=90
「下载文字版」、supp 收录时 role=body 的正文 PDF）。实测 gov 1625 条 / mof 480 条 / supp 1 条。
"""
from __future__ import annotations

import json
import os

import pytest

# 经 tests/conftest.py 唯一引导（门禁 gate_import_bootstrap：tests 层注入只减不增 ⇒ 此处不再自行注入）
from std_lib.scraper_std.attachments import (
    covered_by_body,
    dedupe_attachments,
    match_norm,
    partition_attachments,
    split_record_body_docs,
)


def _body(n: int = 400) -> str:
    return "".join("正文条款第%d条规定了相关事项。" % i for i in range(1, n))


# --------------------------------------------------------------------------- #
# ① 分流判据：内容判据 + 来源角色判据
# --------------------------------------------------------------------------- #
def test_partition_by_content() -> None:
    """附件文本整体已在正文中 ⇒ 判为正文载体；无关内容 ⇒ 判为真实附件。"""
    body = _body(300)
    carrier = {"file_name": "下载Word.docx", "text": body + "尾部说明"}
    real = {"file_name": "附件1-明细表.xlsx", "text": "序号\t名称\t金额\n1\t甲\t100"}
    real_, car = partition_attachments([carrier, real], body)
    assert [a["file_name"] for a in real_] == ["附件1-明细表.xlsx"]
    assert [a["file_name"] for a in car] == ["下载Word.docx"]


def test_partition_by_role_markers() -> None:
    """来源角色判据：mof 的 fileType=90/kind=text_version、站点生成的 html_disguised_doc。"""
    body = "正文完全无关的内容"
    for marker in ({"file_type": "90"}, {"kind": "text_version"},
                   {"attachment_kind": "html_disguised_doc"}, {"role": "body"}):
        entry = {"file_name": "x.doc", "text": "与正文无关的其它文本" * 20, **marker}
        real_, car = partition_attachments([entry], body)
        assert not real_ and len(car) == 1, marker
    plain = {"file_name": "y.pdf", "text": "与正文无关的其它文本" * 20, "attachment_kind": "pdf"}
    assert len(partition_attachments([plain], body)[0]) == 1, "文件类型串不得被误判为载体角色"


def test_covered_by_body_segments() -> None:
    """整体覆盖核对：头部相同但尾部是增量 ⇒ 不应判为"整体覆盖"。"""
    body = _body(200)
    assert covered_by_body(body, match_norm(body))
    assert not covered_by_body(body + "附录：另有若干附表说明" * 30, match_norm(body))


# --------------------------------------------------------------------------- #
# ② 记录级分流：内容守恒 + 幂等 + 聚合字段重算
# --------------------------------------------------------------------------- #
def test_split_record_moves_carriers_and_preserves_content() -> None:
    body = _body(200)
    rec = {
        "full_text": body,
        "attachments": [
            {"file_name": "下载Word.docx", "text": body},                     # 整体覆盖 ⇒ 只留元数据
            {"file_name": "附件1-表.xlsx", "text": "序号\t名称\n1\t甲"},        # 真附件 ⇒ 保留
            {"file_name": "文字版.doc", "kind": "text_version",
             "text": "独立正文（未被网页正文覆盖）" * 20},                      # 角色判据 + 未覆盖 ⇒ 并入正文
        ],
        "attachment_text": "X",
        "attachment_count": 3,
    }
    st = split_record_body_docs(rec)
    assert st["moved"] == 2 and st["kept"] == 1 and st["body_filled"] is True
    assert [a["file_name"] for a in rec["attachments"]] == ["附件1-表.xlsx"]
    assert {c["file_name"] for c in rec["body_docs"]} == {"下载Word.docx", "文字版.doc"}
    assert rec["attachment_count"] == 1
    assert rec["attachment_text"] == "序号\t名称\n1\t甲", "记录级聚合只按**真附件**重算"
    assert "独立正文" in rec["full_text"], "未被正文覆盖的载体文本必须并入正文（内容不丢）"
    assert "text" not in rec["body_docs"][0], "载体条目只留溯源元数据，不重复携带全文"
    # 内容守恒台账（助手内置断言 + 此处显式复核）
    assert st["text_in"] == st["text_kept"] + st["text_to_body"] + st["text_dropped"]


def test_split_record_is_idempotent() -> None:
    body = _body(120)
    rec = {"full_text": body, "attachments": [{"file_name": "下载PDF.pdf", "text": body}],
           "attachment_count": 1}
    split_record_body_docs(rec)
    snap = json.dumps(rec, ensure_ascii=False, sort_keys=True)
    st2 = split_record_body_docs(rec)
    assert st2["moved"] == 0, "二次分流应零改写（幂等）"
    assert json.dumps(rec, ensure_ascii=False, sort_keys=True) == snap


def test_split_noop_without_attachments() -> None:
    rec = {"full_text": "正文"}
    assert split_record_body_docs(rec)["moved"] == 0 and "body_docs" not in rec


# --------------------------------------------------------------------------- #
# ③ XLSX 共享原语：按内容裁剪 + 上限标注（五源统一口径）
# --------------------------------------------------------------------------- #
def _mk_xlsx(tmp_path, *, far_cell: bool):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "序号", "名称"
    ws["A2"], ws["B2"] = 1, "药品甲"
    if far_cell:
        ws.cell(row=3, column=600, value="")     # 空值但存在 ⇒ 声明维度被抬到 600 列
    p = tmp_path / "far.xlsx"
    wb.save(p)
    return str(p)


def test_dedupe_attachments_identical_only() -> None:
    """**完全重复**条目归并（supp 9 条同款）；名称/URL 不同者即使文本相同也保留（保守）。"""
    same = {"name": "关于短期健康保险续保表述的要求", "kind": "本地 .doc 文本提取",
            "content_ref": "data/raw/_sources/x.txt"}
    other = {"name": "另一份附件", "kind": "pdf", "local_path": "data/docs/y.pdf", "text": "同文" * 40}
    # 同名同文但**路径不同** ⇒ 不合并（保守：可能是真实存在的两份同名附件）
    twin = {"name": "另一份附件", "kind": "pdf", "local_path": "data/docs/z.pdf", "text": "同文" * 40}
    out, dup = dedupe_attachments([same, dict(same), dict(same), other, twin])
    assert dup == 2 and len(out) == 3
    # 记录级：仅重复（无载体）时也应回写并重算
    rec = {"full_text": "正文", "attachments": [dict(same), dict(same)], "attachment_count": 2}
    st = split_record_body_docs(rec)
    assert st["deduped"] == 1 and rec["attachment_count"] == 1 and len(rec["attachments"]) == 1
    assert "body_docs" not in rec, "无载体时不应写入空 body_docs"


def test_shared_xlsx_rows_trim_declared_dimension(tmp_path) -> None:
    """声明维度虚高（600 列）时，共享原语只产出**有值内容**行（不得按声明展开）。"""
    from std_lib.scraper_std.crawler_common import cap_xlsx_text, iter_xlsx_text_rows

    rows = [r for r in iter_xlsx_text_rows(open(_mk_xlsx(tmp_path, far_cell=True), "rb").read())]
    text = (chr(10)).join("\t".join(r) for r in rows)
    assert "药品甲" in text and len(text) < 120, "远端空单元格不得撑大文本"
    assert cap_xlsx_text("x" * 100, 10).startswith("x" * 10) and "表格超限" in cap_xlsx_text("x" * 100, 10)
    assert cap_xlsx_text("short", 10) == "short"


def test_collector_extractors_delegate_to_shared_primitives() -> None:
    """**接线守卫**：nfra / pbc 的自建 xlsx 抽取必须引用共享原语（防回退成"按声明维度展开"）。"""
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    for rel, fn in (("modules/regulatory_scrapers/collectors/nfra_attachments_extract.py", "extract_xlsx"),
                    ("modules/regulatory_scrapers/collectors/pbc_parse.py", "extract_xls_text")):
        src = open(os.path.join(root, rel), encoding="utf-8").read()
        assert ("def %s(" % fn) in src, "%s 缺少 %s 定义" % (rel, fn)
        assert "iter_xlsx_text_rows" in src and "cap_xlsx_text" in src, "%s 未使用共享原语" % rel
        assert "openpyxl.load_workbook" not in src and "from openpyxl import" not in src, \
            "%s 仍保留自建 openpyxl 展开（爆量风险）" % rel


@pytest.mark.parametrize("field", ["attachments", "attachment_count"])
def test_write_side_call_sites_wire_split(field: str) -> None:
    """**接线守卫**：gov/mof 写入侧必须调用分流（否则新数据仍会把正文当附件写入）。"""
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    for rel in ("modules/regulatory_scrapers/collectors/gov_collector.py",
                "modules/regulatory_scrapers/collectors/gov_zhengceku.py",
                "modules/regulatory_scrapers/collectors/mof_attachments.py",
                "modules/regulatory_scrapers/collectors/supp_ingest_local_dir.py"):
        src = open(os.path.join(root, rel), encoding="utf-8").read()
        assert "body_docs" in src or "split_record_body_docs" in src, "%s 未接线分流" % rel
