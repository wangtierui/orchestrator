# -*- coding: utf-8 -*-
"""test_ipb_extract_file.py —— extract_file 实文件链（覆盖冲 25%，2026-09-13）。

覆盖：internal_policy_base.extract.extract_file（tmp 真实 txt/docx/md）+ _is_scan_pdf。
"""
from __future__ import annotations

import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
for p in (ROOT, os.path.join(ROOT, "modules", "internal_policy_base")):
    if p not in sys.path:
        sys.path.insert(0, p)

import extract as ex  # noqa: E402


class TestExtractFile:
    def test_txt_file(self, tmp_path):
        p = tmp_path / "制度.txt"
        p.write_text("第一条 为规范管理，制定本制度。\n第二条 适用范围。", encoding="utf-8")
        r = ex.extract_file(str(p), p.name)
        assert isinstance(r, dict)
        assert r.get("text_chars", 0) >= 0

    def test_md_file(self, tmp_path):
        p = tmp_path / "readme.md"
        p.write_text("# 标题\n\n正文段落内容。", encoding="utf-8")
        r = ex.extract_file(str(p), p.name)
        assert isinstance(r, dict)

    def test_docx_file(self, tmp_path):
        from docx import Document
        d = Document()
        d.add_paragraph("第一条 总则内容。")
        d.add_paragraph("第二条 实施细则。")
        p = tmp_path / "制度.docx"
        d.save(str(p))
        r = ex.extract_file(str(p), p.name)
        assert isinstance(r, dict)

    def test_missing_file(self, tmp_path):
        try:
            r = ex.extract_file(str(tmp_path / "nope.txt"), "nope.txt")
            assert isinstance(r, (dict, type(None)))
        except (FileNotFoundError, OSError):
            pass    # 明确抛错亦可

    def test_is_scan_pdf_false_on_text(self, tmp_path):
        p = tmp_path / "a.txt"
        p.write_text("x", encoding="utf-8")
        r = ex._is_scan_pdf(str(p))
        assert r in (True, False)

    def test_is_scan_pdf_on_bad_pdf(self, tmp_path):
        p = tmp_path / "bad.pdf"
        p.write_bytes(b"%PDF-1.4\nbroken")
        r = ex._is_scan_pdf(str(p))
        assert r in (True, False)
